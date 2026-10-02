import streamlit as st
import pandas as pd
from supabase import create_client, Client
from datetime import datetime
import io
import zipfile
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from PIL import Image

# --- PAGE CONFIGURATION ---
st.set_page_config(
    page_title="KEA Comprehensive School Management System",
    page_icon="🏫",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --- SUPABASE CONNECTION SETUP ---
@st.cache_resource
def init_supabase():
    try:
        url = st.secrets["supabase"]["url"]
        key = st.secrets["supabase"]["key"]
        return create_client(url, key)
    except Exception as e:
        st.error(f"Failed to connect to Supabase. Check secrets.toml configuration. Details: {e}")
        return None

supabase = init_supabase()

# =========================================================
# --- SUBSCRIPTION & SYSTEM LICENSE GATEKEEPER ---
# =========================================================
def check_subscription():
    if not supabase:
        return
    try:
        response = (
            supabase.table("school_subscriptions")
            .select("*")
            .eq("school_id", "KEA_COMPREHENSIVE")
            .execute()
        )
        if response.data:
            sub_info = response.data[0]
            expiry_str = sub_info.get("expiry_date")
            status = sub_info.get("status", "INACTIVE").upper()
            
            if expiry_str:
                expiry_date = datetime.strptime(expiry_str, "%Y-%m-%d").date()
                today = datetime.now().date()
                
                # If subscription is inactive or current date has passed the expiry date
                if status != "ACTIVE" or today > expiry_date:
                    st.error("🔒 **System Access Expired**")
                    st.warning(
                        "The subscription license for KEA Comprehensive School SMS has ended or is inactive. "
                        "Please contact the system developer to renew your access."
                    )
                    st.stop()
    except Exception as e:
        pass

check_subscription()

# --- INITIALIZE SESSION STATE ---
if "logged_in" not in st.session_state:
    st.session_state.logged_in = False
    st.session_state.username = ""
    st.session_state.role = ""
    st.session_state.full_name = ""

# --- TEACHER SUBJECT & GRADE ASSIGNMENTS ---
TEACHER_ASSIGNMENTS = {
    "Opondo": {
        "assignments": [
            {"grade": "Grade 7", "subjects": ["ENGLISH", "PRETECHNICAL STUDIES"]},
            {"grade": "Grade 8", "subjects": ["ENGLISH"]},
            {"grade": "Grade 9", "subjects": ["ENGLISH"]}
        ]
    },
    "Lucas": {
        "assignments": [
            {"grade": "Grade 8", "subjects": ["INTEGRATED SCIENCE", "AGRICULTURE"]},
            {"grade": "Grade 9", "subjects": ["MATHEMATICS", "AGRICULTURE"]} # Class Teacher Grade 9
        ]
    },
    "Vincent": {
        "assignments": [
            {"grade": "Grade 7", "subjects": ["SOCIAL STUDIES"]},
            {"grade": "Grade 8", "subjects": ["SOCIAL STUDIES", "KISWAHILI"]}, # Class Teacher Grade 8
            {"grade": "Grade 9", "subjects": ["SOCIAL STUDIES", "KISWAHILI"]}
        ]
    },
    "Grace": {
        "assignments": [
            {"grade": "Grade 7", "subjects": ["MATHEMATICS", "AGRICULTURE"]}, # Class Teacher Grade 7
            {"grade": "Grade 9", "subjects": ["INTEGRATED SCIENCE"]}
        ]
    },
    "Achiyo": {
        "assignments": [
           {"grade": "Grade 9", "subjects": ["CHRISTIAN RELIGIOUS EDUCATION"]}
        ]
    },
    "Balla":{
        "assignments":[
            {"grade": "Grade 8", "subjects": ["CREATIVE ARTS SPORTS"]}
        ]
    },
    "Faith": {
        "assignments": [
            {"grade": "Grade 7", "subjects": ["INTEGRATED SCIENCE", "KISWAHILI"]},
            {"grade": "Grade 8", "subjects": ["CHRISTIAN RELIGIOUS EDUCATION"]}
        ]
    },
    "Elijah": {
        "assignments": [
            {"grade": "Grade 7", "subjects": ["CREATIVE ARTS SPORTS"]},
            {"grade": "Grade 8", "subjects": ["MATHEMATICS", "PRETECHNICAL STUDIES"]},
            {"grade": "Grade 9", "subjects": ["PRETECHNICAL STUDIES", "CREATIVE ARTS SPORTS"]}
        ]
    }
}

# Designated Class Teachers (Allowed to access Results Analysis)
CLASS_TEACHERS = ["Grace", "Vincent", "Lucas"]

LEARNING_AREAS = [
    "MATHEMATICS", "ENGLISH", "KISWAHILI", "INTEGRATED SCIENCE", 
    "AGRICULTURE", "PRETECHNICAL STUDIES", "SOCIAL STUDIES", 
    "CHRISTIAN RELIGIOUS EDUCATION", "CREATIVE ARTS SPORTS"
]

# --- SUBJECT TO POSTGRESQL COLUMN MAPPING HELPER ---
SUBJECT_COLUMN_MAP = {
    "CREATIVE ARTS SPORTS": "creative_arts_sports",
    "CHRISTIAN RELIGIOUS EDUCATION": "christian_religious_education",
    "PRETECHNICAL STUDIES": "pretechnical_studies",
    "INTEGRATED SCIENCE": "integrated_science",
    "SOCIAL STUDIES": "social_studies"
}

def get_db_col_name(subject_name):
    """Converts a subject title to its exact PostgreSQL database column name."""
    subject_clean = subject_name.strip().upper()
    if subject_clean in SUBJECT_COLUMN_MAP:
        return SUBJECT_COLUMN_MAP[subject_clean]
    return subject_name.lower().strip().replace(" ", "_").replace(".", "").replace("(", "").replace(")", "")

# --- AUTHENTICATION & LOGIN SCREEN ---
def login_screen():
    st.markdown("<h1 style='text-align: center; color: #1E3A8A;'>KEA COMPREHENSIVE SCHOOL</h1>", unsafe_allow_html=True)
    st.markdown("<h3 style='text-align: center; color: #4B5563;'>School Management & Information System</h3>", unsafe_allow_html=True)
    
    col1, col2, col3 = st.columns([1, 2, 1])
    with col2:
        if os.path.exists("logo.png"):
            st.image("logo.png", width=150)
        
        st.markdown("### Please Login to Continue")
        username = st.text_input("Username").strip()
        password = st.text_input("Password", type="password").strip()
        
        if st.button("Login", type="primary", use_container_width=True):
            if username == "Admin" and password == "janabi@26!":
                st.session_state.logged_in = True
                st.session_state.username = "Admin"
                st.session_state.role = "DEVELOPER"
                st.session_state.full_name = "System Developer"
                st.success("Welcome back, System Developer!")
                st.rerun()
                
            elif supabase:
                try:
                    res = supabase.table("teachers").select("*").ilike("username", username).execute()
                    if res.data:
                        user = res.data[0]
                        if password == "kea@26" or password == user.get("password"):
                            st.session_state.logged_in = True
                            st.session_state.username = user["username"]
                            st.session_state.role = user["designation"]
                            st.session_state.full_name = user["full_name"]
                            st.success(f"Welcome back, {user['full_name']}!")
                            st.rerun()
                        else:
                            st.error("Incorrect password. Default teacher password is 'kea@26'.")
                    else:
                        st.error("Username not found in teachers directory.")
                except Exception as e:
                    st.error(f"Database error during login: {e}")
            else:
                st.error("Supabase connection not initialized.")

if not st.session_state.logged_in:
    login_screen()
    st.stop()

# --- ROLE & NAVIGATION SETUP ---
role = str(st.session_state.role).upper()
username = st.session_state.username

# Define Permission Roles
is_admin_role = role in ["HOI", "DHOI", "DHO", "SENIOR TEACHER", "DEVELOPER"] or username == "Admin"
is_class_teacher = username in CLASS_TEACHERS or username == "Admin" or role == "DEVELOPER"

st.sidebar.markdown(f"**Logged in as:** {st.session_state.full_name}")
st.sidebar.markdown(f"**Role:** {role}")
if st.sidebar.button("Logout"):
    st.session_state.logged_in = False
    st.rerun()

st.sidebar.markdown("---")
st.sidebar.markdown("### Navigation")

# BUILD DYNAMIC NAVIGATION PERMITTED BY ROLE
nav_options = ["Dashboard"]

if is_admin_role:
    nav_options.append("Students Registration")

nav_options.append("Marks Entry")

if is_class_teacher:
    nav_options.append("Results Analysis")

if is_admin_role:
    nav_options.append("Teachers Portal")

page = st.sidebar.selectbox("Go to", nav_options)

# --- HELPER FUNCTION FOR SUBJECT GRADING ---
def get_subject_performance(score):
    if score is None or pd.isna(score) or float(score) < 1.0:
        return "-", "-", None
    score = float(score)
    if score <= 10: return "BE2", "Below Expectation 2", 1
    elif score <= 20: return "BE1", "Below Expectation 1", 2
    elif score <= 30: return "AE2", "Approaching Expectation 2", 3
    elif score <= 40: return "AE1", "Approaching Expectation 1", 4
    elif score <= 56: return "ME2", "Meeting Expectation 2", 5
    elif score <= 73: return "ME1", "Meeting Expectation 1", 6
    elif score <= 88: return "EE2", "Exceeding Expectation 2", 7
    else: return "EE1", "Exceeding Expectation 1", 8

def calculate_total_grade(total_marks):
    if total_marks <= 112: return "BE2 (Below Expectation 2)"
    elif total_marks <= 224: return "BE1 (Below Expectation 1)"
    elif total_marks <= 336: return "AE2 (Approaching Expectation 2)"
    elif total_marks <= 449: return "AE1 (Approaching Expectation 1)"
    elif total_marks <= 560: return "ME2 (Meeting Expectation 2)"
    elif total_marks <= 672: return "ME1 (Meeting Expectation 1)"
    elif total_marks <= 785: return "EE2 (Exceeding Expectation 2)"
    else: return "EE1 (Exceeding Expectation 1)"

# --- HELPER FUNCTION: FORMULATE TEACHER'S COMMENT ---
def generate_teacher_comment(total_score, top_subject, low_subject):
    overall_grade = calculate_total_grade(total_score)
    if "EE" in overall_grade:
        return f"An excellent performance! Demonstrates outstanding mastery across learning areas, especially in {top_subject}. Keep up the high standard!"
    elif "ME" in overall_grade:
        return f"Good performance. Shows consistent effort and understanding, particularly in {top_subject}. Aim to put more effort into {low_subject} for better results."
    elif "AE" in overall_grade:
        return f"Fair performance. Shows potential in {top_subject}, but requires more focus and practice in {low_subject} to improve overall standing."
    else:
        return f"Below expectations. Needs active academic support and dedicated revision, particularly in {low_subject}. Steady practice in {top_subject} will build confidence."

# =========================================================
# --- PAGE 1: DASHBOARD ---
# =========================================================
if page == "Dashboard":
    st.markdown("<h1 style='text-align: center;'>KEA COMPREHENSIVE SCHOOL</h1>", unsafe_allow_html=True)
    if os.path.exists("logo.png"):
        col_l1, col_l2, col_l3 = st.columns([1, 2, 1])
        with col_l2:
            st.image("logo.png", width=200)
    st.markdown("<h4 style='text-align: center; color: #555;'>Primary (Grade 1-6) & Junior Secondary (Grade 7-9)</h4>", unsafe_allow_html=True)
    st.markdown("---")

    total_students = 0
    total_teachers = 0

    if supabase:
        try:
            students_res = supabase.table("students").select("*", count="exact").execute()
            teachers_res = supabase.table("teachers").select("*", count="exact").execute()
            
            total_students = students_res.count if students_res.count is not None else len(students_res.data)
            total_teachers = teachers_res.count if teachers_res.count is not None else len(teachers_res.data)
        except Exception as e:
            st.error(f"Error fetching dashboard metrics: {e}")

    col1, col2 = st.columns(2)
    col1.metric("Total Students Enrolled", total_students)
    col2.metric("Total Teaching Staff", total_teachers)
    
    st.markdown("---")
    st.info("💡 **Core Focus:** This system is optimized for academic Results Analysis, Student Rosters, and Subject Marks Entry.")

# =========================================================
# --- PAGE 2: STUDENTS REGISTRATION (RESTRICTED ACCESS) ---
# =========================================================
elif page == "Students Registration":
    if not is_admin_role:
        st.error("🔒 Access Denied: Only Developer, HOI, DHO/DHOI, and Senior Teachers can register students.")
        st.stop()
        
    st.header("Students Registration Portal")
    tab_manual, tab_excel, tab_manage = st.tabs(["Manual Registration", "Excel Spreadsheet Upload", "Manage Registered Students"])
    
    with tab_manual:
        with st.form("manual_student_form"):
            adm_no = st.text_input("Admission Number")
            assessment_no = st.text_input("Assessment Number")
            name = st.text_input("Full Name")
            grade = st.selectbox("Grade", [f"Grade {i}" for i in range(1, 10)])
            submit_manual = st.form_submit_button("Admit Student")
            
            if submit_manual:
                if adm_no and name and grade:
                    try:
                        supabase.table("students").insert({
                            "adm_no": adm_no,
                            "assessment_no": assessment_no,
                            "name": name,
                            "grade": grade
                        }).execute()
                        st.success(f"Successfully admitted {name} to {grade}!")
                    except Exception as e:
                        st.error(f"Error saving student: {e}")
                else:
                    st.warning("Please fill in all mandatory fields.")

    with tab_excel:
        st.markdown("Upload an Excel spreadsheet with columns: **Adm No**, **Assessment No**, **Name**, **Grade**")
        uploaded_file = st.file_uploader("Choose Excel file", type=["xlsx", "xls"])
        selected_grade_bulk = st.selectbox("Target Grade for Bulk Upload", [f"Grade {i}" for i in range(1, 10)])
        
        if uploaded_file and st.button("Process & Admit Students"):
            try:
                df = pd.read_excel(uploaded_file)
                df.columns = [str(c).strip().lower() for c in df.columns]
                
                adm_col = next((c for c in df.columns if "adm" in c), None)
                name_col = next((c for c in df.columns if "name" in c), None)
                assess_col = next((c for c in df.columns if "assess" in c), None)
                grade_col = next((c for c in df.columns if "grade" in c), None)
                
                if not adm_col or not name_col:
                    st.error("Excel file must contain at least 'Adm No' and 'Name' columns.")
                    st.stop()

                records = []
                for _, row in df.iterrows():
                    adm_val = str(row[adm_col]).strip() if pd.notna(row[adm_col]) else None
                    name_val = str(row[name_col]).strip() if pd.notna(row[name_col]) else None
                    assess_val = str(row[assess_col]).strip() if assess_col and pd.notna(row[assess_col]) else ""
                    grade_val = str(row[grade_col]).strip() if grade_col and pd.notna(row[grade_col]) else selected_grade_bulk
                    
                    if adm_val and adm_val.lower() != "nan" and name_val and name_val.lower() != "nan":
                        records.append({
                            "adm_no": adm_val,
                            "assessment_no": assess_val,
                            "name": name_val,
                            "grade": grade_val
                        })

                if not records:
                    st.warning("No valid student records found in the uploaded file.")
                else:
                    unique_records_dict = {r["adm_no"]: r for r in records}
                    deduplicated_records = list(unique_records_dict.values())
                    
                    supabase.table("students").upsert(deduplicated_records, on_conflict="adm_no").execute()
                    st.success(f"Successfully uploaded and admitted {len(deduplicated_records)} students!")

            except Exception as e:
                st.error(f"Error processing excel file: {e}")

    with tab_manage:
        st.markdown("### View, Edit & Deregister Students")
        filter_grade = st.selectbox("Filter by Grade", ["All Grades"] + [f"Grade {i}" for i in range(1, 10)])
        
        try:
            query = supabase.table("students").select("*")
            if filter_grade != "All Grades":
                query = query.eq("grade", filter_grade)
            
            res = query.execute()
            students_list = res.data if res.data else []
            
            if not students_list:
                st.info("No students found.")
            else:
                df_students = pd.DataFrame(students_list)
                cols = [c for c in ["adm_no", "assessment_no", "name", "grade"] if c in df_students.columns]
                df_students = df_students[cols]
                
                edited_students_df = st.data_editor(
                    df_students,
                    column_config={
                        "adm_no": st.column_config.TextColumn("Adm No", disabled=True),
                        "assessment_no": st.column_config.TextColumn("Assessment No"),
                        "name": st.column_config.TextColumn("Full Name"),
                        "grade": st.column_config.SelectboxColumn("Grade", options=[f"Grade {i}" for i in range(1, 10)])
                    },
                    use_container_width=True,
                    hide_index=True,
                    key="students_roster_editor"
                )
                
                if st.button("💾 Save Changes to Database", type="primary"):
                    try:
                        records_to_update = edited_students_df.to_dict(orient="records")
                        supabase.table("students").upsert(records_to_update).execute()
                        st.success("Successfully updated student records!")
                        st.rerun()
                    except Exception as e:
                        st.error(f"Error updating records: {e}")

                st.markdown("---")
                col_del1, col_del2 = st.columns([3, 1])
                with col_del1:
                    student_to_delete = st.selectbox(
                        "Select Student to Deregister", 
                        students_list, 
                        format_func=lambda s: f"Adm No: {s['adm_no']} — {s['name']} ({s['grade']})"
                    )
                with col_del2:
                    st.write("")
                    st.write("")
                    confirm_delete = st.button("🚨 Deregister Student", type="secondary")
                    
                if confirm_delete and student_to_delete:
                    adm_del = student_to_delete["adm_no"]
                    name_del = student_to_delete["name"]
                    try:
                        supabase.table("marks").delete().eq("adm_no", adm_del).execute()
                        supabase.table("students").delete().eq("adm_no", adm_del).execute()
                        st.success(f"Successfully deregistered {name_del} (Adm No: {adm_del}).")
                        st.rerun()
                    except Exception as e:
                        st.error(f"Error deregistering student: {e}")

        except Exception as e:
            st.error(f"Error loading students: {e}")

# =========================================================
# --- PAGE 3: MARKS ENTRY (ACCESSIBLE TO ALL TEACHERS) ---
# =========================================================
elif page == "Marks Entry":
    st.header("Subject-Specific Marks Entry Portal")
    
    user_key = st.session_state.username
    teacher_data = TEACHER_ASSIGNMENTS.get(user_key)
    
    if is_admin_role or not teacher_data:
        allowed_grades = [f"Grade {i}" for i in range(1, 10)]
        teacher_grade_map = {g: LEARNING_AREAS for g in allowed_grades}
        st.info("Log-in Mode: Administrator Mode (Full Access)")
    else:
        teacher_grade_map = {}
        for item in teacher_data["assignments"]:
            teacher_grade_map[item["grade"]] = item["subjects"]
        allowed_grades = list(teacher_grade_map.keys())
        st.info(f"Welcome, {st.session_state.full_name}! Select an assigned grade and learning area below.")

    if not allowed_grades:
        st.warning("No grades currently assigned to your profile.")
        st.stop()

    col_g, col_s = st.columns(2)
    with col_g:
        target_grade = st.selectbox("1. Select Grade", allowed_grades)
    
    assigned_subjects = teacher_grade_map.get(target_grade, LEARNING_AREAS)
    
    with col_s:
        target_subject = st.selectbox("2. Select Learning Area", assigned_subjects)

    st.markdown("---")
    st.subheader(f"Marking Sheet: {target_grade} — {target_subject}")
    st.caption("Enter scores out of 100%. Any score less than 1 (x < 1) is saved as Null (-).")

    grade_num = target_grade.replace("Grade", "").strip()
    try:
        students_res = supabase.table("students").select("adm_no, name, grade").execute()
        all_students = students_res.data if students_res.data else []
        students = [
            s for s in all_students 
            if str(s.get("grade", "")).strip().lower() in [target_grade.lower(), grade_num, f"grade {grade_num}"]
        ]
        students = sorted(students, key=lambda x: str(x.get("adm_no", "")))
    except Exception as e:
        students = []
        st.error(f"Error fetching students from database: {e}")

    if not students:
        st.warning(f"No students found under '{target_grade}' in the database.")
    else:
        col_key = get_db_col_name(target_subject)
        existing_marks_map = {}
        try:
            m_res = supabase.table("marks").select(f"adm_no, {col_key}").eq("grade", target_grade).execute()
            if m_res.data:
                for row in m_res.data:
                    existing_marks_map[row["adm_no"]] = row.get(col_key)
        except Exception as e:
            st.caption(f"Note: Could not load prior marks automatically: {e}")

        editor_data = []
        for s in students:
            adm = s["adm_no"]
            val = existing_marks_map.get(adm)
            score_val = float(val) if val is not None and float(val) >= 1.0 else 0.0
            
            editor_data.append({
                "Adm No": adm,
                "Student Name": s["name"],
                "Score (%)": score_val
            })

        df_editor = pd.DataFrame(editor_data)

        tab_manual_marks, tab_excel_marks = st.tabs(["Manual Sheet Input", "Upload Excel Marks"])

        with tab_manual_marks:
            with st.form(key=f"form_{target_grade}_{col_key}"):
                st.markdown("### Class List Marks Input")
                edited_df = st.data_editor(
                    df_editor,
                    column_config={
                        "Adm No": st.column_config.TextColumn("Adm No", disabled=True),
                        "Student Name": st.column_config.TextColumn("Student Name", disabled=True),
                        "Score (%)": st.column_config.NumberColumn(
                            "Score (%)", 
                            min_value=0.0, 
                            max_value=100.0, 
                            step=1.0,
                            help="Enter score between 1 and 100. Scores < 1 are saved as Null (-)"
                        )
                    },
                    use_container_width=True,
                    hide_index=True
                )
                
                preview_submitted = st.form_submit_button("Preview Marks Sheet", type="secondary")

            if preview_submitted or st.session_state.get(f"preview_active_{target_grade}_{col_key}", False):
                st.session_state[f"preview_active_{target_grade}_{col_key}"] = True
                
                st.markdown("---")
                st.markdown("### 📋 Preview Marks & Performance Levels")
                
                preview_rows = []
                for _, row in edited_df.iterrows():
                    raw_score = row["Score (%)"]
                    grade_code, perf_level, pts = get_subject_performance(raw_score)
                    
                    preview_rows.append({
                        "Adm No": row["Adm No"],
                        "Student Name": row["Student Name"],
                        "Score (%)": raw_score if raw_score >= 1.0 else "-",
                        "Grade": grade_code,
                        "Performance Level": perf_level
                    })
                    
                preview_table = pd.DataFrame(preview_rows)
                st.dataframe(preview_table, use_container_width=True)

                if st.button("✅ Confirm & Submit to Results Analysis", type="primary"):
                    records_to_upsert = []
                    for _, row in edited_df.iterrows():
                        adm = str(row["Adm No"])
                        raw_score = row["Score (%)"]
                        
                        existing = {}
                        try:
                            ex_res = supabase.table("marks").select("*").eq("adm_no", adm).execute()
                            if ex_res.data:
                                existing = ex_res.data[0]
                        except:
                            pass

                        rec = existing if existing else {"adm_no": adm, "grade": target_grade}
                        rec[col_key] = float(raw_score) if raw_score >= 1.0 else None
                        
                        total = 0.0
                        for sub in LEARNING_AREAS:
                            s_key = get_db_col_name(sub)
                            val = rec.get(s_key)
                            if val is not None and float(val) >= 1.0:
                                total += float(val)
                        
                        rec["total_marks"] = total
                        records_to_upsert.append(rec)

                    try:
                        supabase.table("marks").upsert(records_to_upsert, on_conflict="adm_no").execute()
                        st.success(f"🎉 Successfully updated {target_subject} marks for {len(records_to_upsert)} students in {target_grade}!")
                        st.session_state[f"preview_active_{target_grade}_{col_key}"] = False
                    except Exception as e:
                        st.error(f"Failed to submit marks to database: {e}")

        with tab_excel_marks:
            st.markdown(f"Upload an Excel sheet for **{target_grade} — {target_subject}**.")
            marks_file = st.file_uploader("Choose Excel File", type=["xlsx", "xls"], key=f"uploader_{target_grade}_{col_key}")
            
            if marks_file and st.button("Upload Excel Marks Sheet"):
                try:
                    df_upload = pd.read_excel(marks_file)
                    records_to_upsert = []
                    
                    score_col = None
                    for c in df_upload.columns:
                        if c.lower() in ["score", "marks", target_subject.lower()]:
                            score_col = c
                            break
                    
                    if not score_col and len(df_upload.columns) >= 2:
                        score_col = df_upload.columns[1]

                    for _, row in df_upload.iterrows():
                        adm = str(row["Adm No"]).strip()
                        raw_val = row[score_col]
                        try:
                            score_num = float(raw_val)
                        except:
                            score_num = 0.0

                        existing = {}
                        try:
                            ex_res = supabase.table("marks").select("*").eq("adm_no", adm).execute()
                            if ex_res.data:
                                existing = ex_res.data[0]
                        except:
                            pass

                        rec = existing if existing else {"adm_no": adm, "grade": target_grade}
                        rec[col_key] = score_num if score_num >= 1.0 else None
                        
                        total = 0.0
                        for sub in LEARNING_AREAS:
                            s_key = get_db_col_name(sub)
                            val = rec.get(s_key)
                            if val is not None and float(val) >= 1.0:
                                total += float(val)
                        rec["total_marks"] = total
                        records_to_upsert.append(rec)

                    supabase.table("marks").upsert(records_to_upsert, on_conflict="adm_no").execute()
                    st.success(f"Successfully processed and uploaded marks for {len(records_to_upsert)} students!")
                except Exception as e:
                    st.error(f"Error processing uploaded file: {e}")

# =========================================================
# --- PAGE 4: RESULTS ANALYSIS (RESTRICTED TO CLASS TEACHERS & DEVELOPER) ---
# =========================================================
elif page == "Results Analysis":
    if not is_class_teacher:
        st.error("🔒 Access Denied: Results Analysis is strictly reserved for Class Teachers and the Developer.")
        st.stop()
        
    st.header("Results Analysis Hub & Report Forms")
    
    analysis_grade = st.selectbox("Select Grade for Analysis", [f"Grade {i}" for i in range(1, 10)], key="analysis_grade_select")
    
    tab_overview, tab_reports = st.tabs(["Tab 1: General Performance Overview", "Tab 2: Assessment Report Forms Hub"])
    
    grade_num = analysis_grade.replace("Grade", "").strip()
    try:
        m_res = supabase.table("marks").select("*").execute()
        all_marks = m_res.data if m_res.data else []
        marks_data = [
            m for m in all_marks 
            if str(m.get("grade", "")).strip().lower() in [analysis_grade.lower(), grade_num, f"grade {grade_num}"]
        ]
        
        s_res = supabase.table("students").select("*").execute()
        students_dict = {}
        if s_res.data:
            for s in s_res.data:
                clean_adm = str(s.get("adm_no", "")).strip()
                students_dict[clean_adm] = s
    except Exception as e:
        marks_data, students_dict = [], {}
        st.error(f"Error fetching analysis data: {e}")

    def process_student_scores(m_row):
        total = 0.0
        valid_count = 0
        for sub in LEARNING_AREAS:
            k = get_db_col_name(sub)
            val = m_row.get(k)
            if val is not None and not pd.isna(val) and float(val) >= 1.0:
                total += float(val)
                valid_count += 1
        return total, valid_count

    active_ranked_students = []
    for m in marks_data:
        adm_str = str(m.get("adm_no", "")).strip()
        s_name = students_dict.get(adm_str, {}).get("name", f"Student {adm_str}")
        tot_marks, valid_cnt = process_student_scores(m)
        
        if valid_cnt > 0:
            active_ranked_students.append({
                "adm_no": adm_str,
                "name": s_name,
                "total_marks": tot_marks,
                "marks_row": m
            })

    active_ranked_students = sorted(active_ranked_students, key=lambda x: x["total_marks"], reverse=True)
    for idx, item in enumerate(active_ranked_students):
        item["rank"] = idx + 1

    # TAB 1: OVERVIEW
    with tab_overview:
        if not active_ranked_students:
            st.info(f"No active student records with submitted marks found in {analysis_grade}.")
        else:
            st.markdown("### 🏆 Top Ten Students")
            top_10_rows = []
            for item in active_ranked_students[:10]:
                top_10_rows.append({
                    "Rank": item["rank"],
                    "Adm No": item["adm_no"],
                    "Student Name": item["name"],
                    "Total Marks": item["total_marks"],
                    "Overall Grade": calculate_total_grade(item["total_marks"])
                })
            st.dataframe(pd.DataFrame(top_10_rows), use_container_width=True)
            
            st.markdown("---")

            subject_stats = []
            df_marks = pd.DataFrame(marks_data)
            
            for sub in LEARNING_AREAS:
                col_name = get_db_col_name(sub)
                if col_name in df_marks.columns:
                    valid_scores = [float(v) for v in df_marks[col_name] if v is not None and not pd.isna(v) and float(v) >= 1.0]
                    valid_count = len(valid_scores)
                    
                    if valid_count > 0:
                        total_sub_marks = sum(valid_scores)
                        mean_score = total_sub_marks / valid_count
                        total_points = sum([get_subject_performance(s)[2] for s in valid_scores if get_subject_performance(s)[2] is not None])
                        mean_points = total_points / valid_count
                    else:
                        total_sub_marks, mean_score, mean_points = 0.0, 0.0, 0.0

                    subject_stats.append({
                        "Learning Area": sub,
                        "Valid Students": valid_count,
                        "Total Marks": round(total_sub_marks, 1),
                        "Mean Score (%)": round(mean_score, 2),
                        "Aggregate Mean Points": round(mean_points, 2)
                    })

            df_subject_stats = pd.DataFrame(subject_stats)

            if not df_subject_stats.empty and df_subject_stats["Mean Score (%)"].max() > 0:
                best_subject = df_subject_stats.sort_values(by="Mean Score (%)", ascending=False).iloc[0]
                st.markdown(f"### ⭐ Best Performed Learning Area")
                st.success(
                    f"**{best_subject['Learning Area']}** with a Mean Score of **{best_subject['Mean Score (%)']}%** "
                    f"and Aggregate Points of **{best_subject['Aggregate Mean Points']}** (based on {best_subject['Valid Students']} valid submissions)."
                )
            
            st.markdown("---")
            st.markdown("### 📊 Performance Summary for Each Learning Area")
            st.dataframe(df_subject_stats, use_container_width=True)

    # TAB 2: REPORTS HUB
    with tab_reports:
        st.markdown("### Assessment Report Form Generator")
        
        col_s1, col_s2 = st.columns(2)
        with col_s1:
            term_val = st.selectbox("Term", ["Term 1", "Term 2", "Term 3"], key="rep_term")
            opening_date = st.text_input("Opening Date", "14/09/2026", key="rep_open")
        with col_s2:
            closing_date = st.text_input("Closing Date", "04/12/2026", key="rep_close")
            stamp_upload = st.file_uploader("Upload Stamp Image (stamp.png)", type=["png", "jpg"], key="rep_stamp")

        st.markdown("---")
        
        if not active_ranked_students:
            st.info(f"No students with active exam scores found in {analysis_grade}.")
        else:
            def generate_student_docx_object(student_obj):
                adm_str = student_obj["adm_no"]
                student_name = student_obj["name"]
                s_row = student_obj["marks_row"]
                student_rank = student_obj["rank"]

                if os.path.exists("report.docx"):
                    doc = Document("report.docx")
                else:
                    doc = Document()

                total_score = 0.0
                top_sub, low_sub = "Mathematics", "English"
                best_s, worst_s = -1, 101

                subject_scores = {}
                for sub in LEARNING_AREAS:
                    col_key = get_db_col_name(sub)
                    val = s_row.get(col_key)
                    if val is not None and not pd.isna(val) and float(val) >= 1.0:
                        s_num = float(val)
                        total_score += s_num
                        if s_num > best_s: best_s = s_num; top_sub = sub
                        if s_num < worst_s: worst_s = s_num; low_sub = sub
                        g_code, p_lvl, pts = get_subject_performance(s_num)
                        subject_scores[sub.upper()] = (str(s_num), p_lvl)
                    else:
                        subject_scores[sub.upper()] = ("-", "-")

                overall_g = calculate_total_grade(total_score)
                teacher_comment = generate_teacher_comment(total_score, top_sub, low_sub)

                for t in doc.tables:
                    for row in t.rows:
                        row_text = " ".join([cell.text for cell in row.cells]).upper()
                        
                        for cell in row.cells:
                            if "NAME:" in cell.text.upper():
                                cell.text = f"NAME: {student_name}"
                            elif "ADM NO:" in cell.text.upper() or "ADM.NO:" in cell.text.upper():
                                cell.text = f"ADM NO: {adm_str}"
                            elif "TERM:" in cell.text.upper():
                                cell.text = f"TERM: {term_val}"
                            elif "POSITION:" in cell.text.upper() or "RANK:" in cell.text.upper():
                                cell.text = f"POSITION: #{student_rank}"
                            elif "YEAR:" in cell.text.upper():
                                cell.text = f"YEAR: 2026"
                                
                        for sub_name, (score_str, perf_str) in subject_scores.items():
                            if sub_name in row_text and len(row.cells) >= 5:
                                row.cells[3].text = score_str
                                row.cells[4].text = perf_str

                replacements = {
                    "TOTAL MARKS: _____": f"TOTAL MARKS: {total_score:.0f}",
                    "GENERAL PERFORMANCE LEVEL:": f"GENERAL PERFORMANCE LEVEL: {overall_g}",
                    "TEACHERS GENERAL COMMENT:": f"TEACHERS GENERAL COMMENT:\n\"{teacher_comment}\"",
                    "CLOSING DATE: __________": f"CLOSING DATE: {closing_date}",
                    "OPENING DATE: __________": f"OPENING DATE: {opening_date}",
                }

                for p in doc.paragraphs:
                    for k, v in replacements.items():
                        if k in p.text:
                            p.text = p.text.replace(k, v)

                if stamp_upload or os.path.exists("stamp.png"):
                    stamp_source = stamp_upload if stamp_upload else "stamp.png"
                    for p in doc.paragraphs:
                        if "H.O.I STAMP" in p.text or "{{STAMP}}" in p.text:
                            run = p.add_run()
                            run.add_picture(stamp_source, width=Inches(1.2))
                            break

                return doc

            st.subheader("1. Live Student Report Form Preview")
            
            selected_student_obj = st.selectbox(
                "Select Student for Live Report Form Preview", 
                active_ranked_students, 
                format_func=lambda x: f"Rank #{x['rank']} | Adm No: {x['adm_no']} | Name: {x['name']} (Total: {x['total_marks']:.0f})"
            )

            if selected_student_obj:
                doc_obj = generate_student_docx_object(selected_student_obj)
                
                s_name = selected_student_obj['name']
                s_adm = selected_student_obj['adm_no']
                s_rank = selected_student_obj['rank']
                s_total = selected_student_obj['total_marks']
                s_row = selected_student_obj['marks_row']
                overall_g = calculate_total_grade(s_total)
                
                top_s, low_s = "Mathematics", "English"
                b_val, w_val = -1, 101
                for sub in LEARNING_AREAS:
                    col_key = get_db_col_name(sub)
                    val = s_row.get(col_key)
                    if val is not None and not pd.isna(val) and float(val) >= 1.0:
                        v = float(val)
                        if v > b_val: b_val = v; top_s = sub
                        if v < w_val: w_val = v; low_s = sub
                teacher_comment = generate_teacher_comment(s_total, top_s, low_s)

                html_preview = f"""
                <div style="border: 2px solid #374151; padding: 25px; border-radius: 8px; background-color: #FFFFFF; color: #111827; font-family: Arial, sans-serif;">
                    <div style="text-align: center; font-weight: bold; margin-bottom: 15px;">
                        <h2 style="margin: 0; color: #1E3A8A;">KEA COMPREHENSIVE SCHOOL</h2>
                        <p style="margin: 2px 0;">P.O. BOX 557-40400 SUNA MIGORI</p>
                        <h4 style="margin: 5px 0; text-decoration: underline;">STUDENT ASSESSMENT REPORT</h4>
                    </div>

                    <table style="width: 100%; border-collapse: collapse; margin-bottom: 15px; font-size: 13px; border: 1px solid #374151;">
                        <tr>
                            <td style="border: 1px solid #374151; padding: 6px; font-weight: bold;">NAME: {s_name}</td>
                            <td style="border: 1px solid #374151; padding: 6px; font-weight: bold;">YEAR: 2026</td>
                        </tr>
                        <tr>
                            <td style="border: 1px solid #374151; padding: 6px; font-weight: bold;">TERM: {term_val}</td>
                            <td style="border: 1px solid #374151; padding: 6px; font-weight: bold;">POSITION: #{s_rank}</td>
                        </tr>
                        <tr>
                            <td style="border: 1px solid #374151; padding: 6px; font-weight: bold;">ADM NO: {s_adm}</td>
                            <td style="border: 1px solid #374151; padding: 6px; font-weight: bold;">GRADE: {analysis_grade}</td>
                        </tr>
                    </table>

                    <table style="width: 100%; border-collapse: collapse; font-size: 12px; border: 1px solid #374151;">
                        <tr style="background-color: #F3F4F6; font-weight: bold; text-align: left;">
                            <th style="border: 1px solid #374151; padding: 6px;">S/N</th>
                            <th style="border: 1px solid #374151; padding: 6px;">CODE</th>
                            <th style="border: 1px solid #374151; padding: 6px;">LEARNING AREA</th>
                            <th style="border: 1px solid #374151; padding: 6px;">MARKS SCORED</th>
                            <th style="border: 1px solid #374151; padding: 6px;">PERFORMANCE LEVEL</th>
                        </tr>
                """

                for idx, sub in enumerate(LEARNING_AREAS, 1):
                    col_key = get_db_col_name(sub)
                    val = s_row.get(col_key)
                    code_val = f"90{idx}" if idx < 10 else f"9{idx}"
                    
                    if val is not None and not pd.isna(val) and float(val) >= 1.0:
                        s_num = float(val)
                        g_code, p_lvl, pts = get_subject_performance(s_num)
                        m_str = f"{s_num:.0f}"
                        p_str = p_lvl
                    else:
                        m_str = "-"
                        p_str = "-"

                    html_preview += f"""
                        <tr>
                            <td style="border: 1px solid #374151; padding: 6px;">{idx}</td>
                            <td style="border: 1px solid #374151; padding: 6px;">{code_val}</td>
                            <td style="border: 1px solid #374151; padding: 6px; font-weight: bold;">{sub.upper()}</td>
                            <td style="border: 1px solid #374151; padding: 6px; font-weight: bold; color: #1E3A8A;">{m_str}</td>
                            <td style="border: 1px solid #374151; padding: 6px;">{p_str}</td>
                        </tr>
                    """

                html_preview += f"""
                    </table>

                    <div style="margin-top: 15px; font-size: 13px; border-top: 2px solid #374151; padding-top: 10px;">
                        <p style="margin: 5px 0;"><strong>TOTAL MARKS:</strong> <span style="font-size: 15px; color: #1E3A8A;">{s_total:.0f}</span> &nbsp;|&nbsp; <strong>GENERAL PERFORMANCE LEVEL:</strong> {overall_g}</p>
                        <p style="margin: 10px 0 3px 0;"><strong>TEACHER'S GENERAL COMMENT:</strong></p>
                        <div style="border-bottom: 1px dashed #374151; padding: 4px; font-style: italic; background-color: #F9FAFB;">"{teacher_comment}"</div>
                        
                        <div style="display: flex; justify-content: space-between; margin-top: 15px;">
                            <div><strong>CLOSING DATE:</strong> {closing_date}</div>
                            <div><strong>OPENING DATE:</strong> {opening_date}</div>
                        </div>

                        <div style="display: flex; justify-content: space-between; margin-top: 20px; font-weight: bold;">
                            <div>CLASS TEACHER SIGNATURE: ____________</div>
                            <div>H.O.I STAMP & SIGNATURE: ____________</div>
                        </div>
                    </div>

                </div>
                """

                st.components.v1.html(html_preview, height=620, scrolling=True)

                out_stream = io.BytesIO()
                doc_obj.save(out_stream)
                out_stream.seek(0)
                
                st.download_button(
                    label=f"📥 Download Assessment Report Form for {s_name} (.docx)",
                    data=out_stream,
                    file_name=f"Rank_{s_rank}_{s_adm}_{s_name.replace(' ', '_')}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    type="primary"
                )

            st.markdown("---")

            st.subheader("2. Ranked Batch Reports Generator (Entire Grade)")
            if st.button(f"📦 Generate All Ranked Reports for {analysis_grade} (.zip)", type="secondary"):
                with st.spinner("Compiling ranked report forms... Please wait."):
                    zip_buffer = io.BytesIO()
                    
                    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zip_file:
                        for s_item in active_ranked_students:
                            single_doc = generate_student_docx_object(s_item)
                            
                            doc_bytes = io.BytesIO()
                            single_doc.save(doc_bytes)
                            doc_bytes.seek(0)
                            
                            file_filename = f"Rank_{s_item['rank']:02d}_{s_item['adm_no']}_{s_item['name'].replace(' ', '_')}_Report.docx"
                            zip_file.writestr(file_filename, doc_bytes.getvalue())

                    zip_buffer.seek(0)
                    st.success(f"Successfully compiled {len(active_ranked_students)} ranked report forms into ZIP archive!")
                    
                    st.download_button(
                        label=f"📥 Download All {len(active_ranked_students)} Ranked Reports (.zip)",
                        data=zip_buffer,
                        file_name=f"{analysis_grade.replace(' ', '_')}_Ranked_Assessment_Reports.zip",
                        mime="application/zip"
                    )

# =========================================================
# --- PAGE 5: TEACHERS PORTAL (RESTRICTED ACCESS) ---
# =========================================================
elif page == "Teachers Portal":
    if not is_admin_role:
        st.error("🔒 Access Denied: Teachers Portal directory access is restricted to Developer, HOI, DHO/DHOI, and Senior Teachers.")
        st.stop()
        
    st.header(f"Teachers Portal — {st.session_state.full_name}")
    
    user_key = st.session_state.username
    teacher_info = TEACHER_ASSIGNMENTS.get(user_key)
    
    st.markdown("### School Administrative Overview")
    try:
        t_res = supabase.table("teachers").select("full_name", "username", "designation").execute()
        if t_res.data:
            st.dataframe(pd.DataFrame(t_res.data), use_container_width=True)
    except Exception as e:
        st.error(f"Error loading teachers list: {e}")

    if teacher_info:
        st.markdown("---")
        st.markdown("### Your Assigned Learning Areas & Responsibilities")
        for item in teacher_info["assignments"]:
            grade = item["grade"]
            subs = ", ".join(item["subjects"])
            st.success(f"📌 **{grade}:** Teaching **{subs}**")
            
            try:
                st_res = supabase.table("students").select("adm_no", "assessment_no", "name").eq("grade", grade).execute()
                if st_res.data:
                    st.markdown(f"**Registered Students in {grade}:**")
                    st.dataframe(pd.DataFrame(st_res.data), use_container_width=True)
                else:
                    st.caption(f"No students currently registered under {grade}.")
            except Exception as e:
                st.error(f"Could not load students for {grade}: {e}")